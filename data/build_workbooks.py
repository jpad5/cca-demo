"""Generate synthetic Zava Insurance contents-inventory workbooks for the demo.

Each workbook has one sheet, "Inventory", with two Excel tables that the
CAT Contents Intake workflow reads with Excel Online (Business):

  ClaimHeader  one row: Claim Number, Insured Name, Policy Number, Date of Loss, Province, CAT Code
  Inventory    Ref, Room, Description, Brand, Model, Supplier, Qty, Years, Months, Original Cost, Notes

All names and data are fictional.

Usage:
  pip install openpyxl
  python data/build_workbooks.py                                       # every claim in data/claims.csv
  python data/build_workbooks.py --claim ZAV-1050 --template ZAV-1043  # one extra rehearsal claim
"""
import argparse
import csv
import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = Path(__file__).resolve().parent
HEADER_COLS = ["Claim Number", "Insured Name", "Policy Number", "Date of Loss", "Province", "CAT Code"]
ITEM_COLS = ["Ref", "Room", "Description", "Brand", "Model", "Supplier", "Qty", "Years", "Months", "Original Cost", "Notes"]
NUMERIC = {"Ref": int, "Qty": int, "Years": int, "Months": int, "Original Cost": float}
WIDTHS = {"A": 14, "B": 16, "C": 42, "D": 15, "E": 17, "F": 12, "G": 7, "H": 8, "I": 9, "J": 14, "K": 32}


def load_claims():
    with open(HERE / "claims.csv", newline="", encoding="utf-8") as f:
        return {row["claim_number"]: row for row in csv.DictReader(f)}


def load_items(template):
    with open(HERE / "inventories" / f"{template}.csv", newline="", encoding="utf-8") as f:
        items = []
        for row in csv.DictReader(f):
            item = {}
            for col in ITEM_COLS:
                value = (row.get(col) or "").strip()
                item[col] = NUMERIC[col](value) if value and col in NUMERIC else (value or None)
            items.append(item)
        return items


def build(claim, items, out_dir):
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventory"

    ws["A1"] = "Zava Insurance · Contents Inventory"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = "List each damaged or lost item on its own row. Leave unknown values blank."
    ws["A2"].font = Font(italic=True, color="666666")

    # Claim header table (rows 4-5)
    for j, col in enumerate(HEADER_COLS, start=1):
        ws.cell(row=4, column=j, value=col)
    loss_date = dt.date.fromisoformat(claim["date_of_loss"])
    values = [claim["claim_number"], claim["insured_name"], claim["policy_number"],
              loss_date, claim["province"], claim["cat_code"]]
    for j, value in enumerate(values, start=1):
        ws.cell(row=5, column=j, value=value)
    ws.cell(row=5, column=4).number_format = "yyyy-mm-dd"
    header_table = Table(displayName="ClaimHeader", ref="A4:F5")
    header_table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=False)
    ws.add_table(header_table)

    # Inventory table (header on row 7)
    start = 7
    for j, col in enumerate(ITEM_COLS, start=1):
        ws.cell(row=start, column=j, value=col)
    for i, item in enumerate(items, start=start + 1):
        for j, col in enumerate(ITEM_COLS, start=1):
            ws.cell(row=i, column=j, value=item[col])
        ws.cell(row=i, column=ITEM_COLS.index("Original Cost") + 1).number_format = "#,##0.00"
        ws.cell(row=i, column=ITEM_COLS.index("Description") + 1).alignment = Alignment(wrap_text=True)
    end_row = start + len(items)
    inventory_table = Table(displayName="Inventory", ref=f"A{start}:{get_column_letter(len(ITEM_COLS))}{end_row}")
    inventory_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(inventory_table)

    for column, width in WIDTHS.items():
        ws.column_dimensions[column].width = width
    ws.freeze_panes = f"A{start + 1}"

    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"{claim['claim_number']}.xlsx"
    wb.save(path)
    return path


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--claim", help="Build one claim. If it isn't in claims.csv, rehearsal header values are used.")
    parser.add_argument("--template", help="Inventory template to use with --claim, for example ZAV-1043.")
    parser.add_argument("--out", default=str(HERE / "workbooks"), help="Output folder (default: data/workbooks)")
    args = parser.parse_args()

    claims = load_claims()
    out_dir = Path(args.out)
    if args.claim:
        claim = dict(claims.get(args.claim) or {
            "claim_number": args.claim, "insured_name": "Rehearsal Claim",
            "policy_number": f"ZHP-REH-{args.claim.split('-')[-1]}", "date_of_loss": "2026-08-14",
            "province": "BC", "cat_code": "CAT-26-07", "inventory_template": "ZAV-1043"})
        template = args.template or claim["inventory_template"]
        targets = [(claim, template)]
    else:
        targets = [(c, c["inventory_template"]) for c in claims.values()]

    for claim, template in targets:
        path = build(claim, load_items(template), out_dir)
        print(f"{path}  ({template} template)")


if __name__ == "__main__":
    main()
